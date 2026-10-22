#!/usr/bin/env python3
"""Tests for the Part 3 leftovers (2026-09-30).

    python -m harness.test_leftovers

Each is a small piece with no Word in it: brackets turned round a number or
an address, a table's caption row read again under it, page numbers in
stretches, notes numbered plainly, a note in two paragraphs, signs from a
font of signs, web addresses as links, and Word's line numbers.
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from convert.footnotes import NOTE_PARA, build_xml
from convert.model import Block, PARAGRAPH, FOOTNOTE
from convert.reader import _held_pieces, _split_typed_notes, _unturn
from convert.writer import _link_runs, _page_stretches, _symbol_runs

FAILED = []


def check(name, got, want):
    ok = got == want
    print(f"  {'ok  ' if ok else 'FAIL'}  {name:<56} got {got!r}")
    if not ok:
        print(f"        want {want!r}")
        FAILED.append(name)


def main():
    print("brackets turned round one piece")
    check("a caption's number", _unturn("الشكل رقم )59( :"), "الشكل رقم (59) :")
    check("an address", _unturn("داخل مصر)https://www.total.eg( ."),
          "داخل مصر(https://www.total.eg) .")
    check("an address between two openings",
          _unturn("∗ (www.globalreporting.org( هي"), "∗ (www.globalreporting.org) هي")
    check("a year in Arabic digits", _unturn("إليه )٢٠١٩( al"), "إليه (٢٠١٩) al")
    check("a formula stays", _unturn("وفق Y)-(T k)"), "وفق Y)-(T k)")
    check("joined to a letter stays", _unturn(" )i(c في"), " )i(c في")
    check("the right way round stays", _unturn("(59) و (www.a.com)"), "(59) و (www.a.com)")

    print("\na table's caption row read again under it")
    cells = {"M𝑜𝑑𝑒", "4:", "𝑓=", "6.08", "𝐻𝑧", "𝑀𝑜𝑑𝑒", "5:", "6.25", "6:"}
    again = "M 𝑜𝑑𝑒 4: 𝑓= 6.08 𝐻𝑧 𝑀𝑜𝑑𝑒6: 𝑓= 6.25 𝐻𝑧 ا".split()
    check("a word in two pieces, two drawn together, a stray letter",
          _held_pieces(again, cells), True)
    check("a word the table lacks", _held_pieces(again + ["الشكل"], cells), False)
    check("its leading words", _held_pieces(again + ["الشكل", "(17)"], cells, lead=True),
          len(again))

    print("\npage numbers in stretches")

    def foot(page, text):
        return Block(PARAGRAPH, text, page=page)
    cut = [foot(p, str(n)) for p, n in zip(range(13), [1, 2, 3, 21, 22, 23, 24, 25,
                                                      26, 27, 28, 54, 55])]
    check("1-3, 21-28, 54-55", _page_stretches(cut), ((0, 1), {3: 18, 11: 43}))
    check("a number going down is not",
          _page_stretches([foot(p, str(n)) for p, n in enumerate([5, 6, 7, 1, 2, 3])]),
          (None, None))

    print("\nnotes numbered plainly")
    body = Block(PARAGRAPH, "نص¹ ونص² ونص³".replace("¹", "(1)").replace("²", "(2)")
                 .replace("³", "(3)"), page=4)
    note = Block(FOOTNOTE, "1 والتون مايكل (1998) ص 35 . 2 سامية احمد ص180 . 3 نفس المرجع ص185 .",
                 page=4)
    got = [b.text for b in _split_typed_notes([body, note]) if b.kind == FOOTNOTE]
    check("three notes", got, ["1 والتون مايكل (1998) ص 35 .", "2 سامية احمد ص180 .",
                               "3 نفس المرجع ص185 ."])
    verse = Block(FOOTNOTE, "(8) سورة الحجر: (9).", page=4)
    body9 = Block(PARAGRAPH, "نص (8) ونص (9)", page=4)
    check("a verse number is not the next note",
          [b.text for b in _split_typed_notes([body9, verse]) if b.kind == FOOTNOTE],
          ["(8) سورة الحجر: (9)."])

    print("\na note in two paragraphs")
    xml = build_xml([(2, f"سورة النساء: (163). {NOTE_PARA} انظر: الجزائري.")])
    note_xml = xml[xml.index('w:id="2"'):]
    check("two paragraphs in the note", note_xml.count("<w:p>"), 2)
    check("the mark itself is not written", NOTE_PARA in xml, False)

    print("\nsigns from a font of signs, and links")
    doc = Document()
    p = doc.add_paragraph()
    p.add_run(" الباب الثاني see www.example.org.")
    _symbol_runs(doc.element.body, {"": "Webdings"})
    _link_runs(doc)
    body_xml = doc.element.body.xml
    check("the sign in Webdings", bool(re.search(r'w:ascii="Webdings"[^>]*/>\s*</w:rPr>\s*'
                                                 r'<w:t[^>]*></w:t>', body_xml)), True)
    check("the address a link", "<w:hyperlink" in body_xml, True)
    check("the full stop left out of it",
          bool(re.search(r"<w:t[^>]*>www\.example\.org</w:t>", body_xml)), True)
    check("the text kept", "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body_xml)),
          " الباب الثاني see www.example.org.")

    print("\nfound by the held-out theses (1 Oct)")
    from convert.pdftext import apply_fixes
    turned = {"هلل": "لله"}
    check("a turned word with و joined in front", apply_fixes("َّ وَهلل َّ", turned),
          "َّ وَلله َّ")
    check("with ف", apply_fixes("فهلل الأمر", turned), "فلله الأمر")
    check("a longer word that ends in it stays", apply_fixes("كتابهلل", turned),
          "كتابهلل")
    from convert.order import _as_stored
    placed = [{"c": c, "bbox": (x, 362, x + 4, 386)}
              for c, x in ((" ", 50.2), ("ع", 54.1), ("و", 59.5))]
    stored = {(" ", 50.2, 362): [1], ("و", 59.5, 362): [2], ("ع", 54.1, 362): [3]}
    check("as stored: space, و, ع", _as_stored(placed, list(" عو"), stored),
          [0, 2, 1])
    check("in stored order already: nothing",
          _as_stored([placed[0], placed[2], placed[1]], list(" وع"), stored), None)

    print("\nfound by the made-up plain set (22 Oct)")
    from convert.order import _order_mixed, _reorder, _stray
    # وفي عام2015 ... لخا م: the space drawn between عام and 2015, stored
    # between the ا and م of لخام (plain008 page 3).
    ln = [("ل", 365.0, 367.7), ("خ", 358.1, 365.0), ("ا", 355.2, 358.1),
          (" ", 206.6, 210.7), ("م", 350.5, 354.9), (" ", 346.6, 350.7),
          ("ع", 217.5, 224.3), ("ا", 214.5, 217.5), ("م", 210.1, 214.4),
          ("2", 179.2, 186.0), ("0", 186.0, 192.8), ("1", 192.9, 199.7),
          ("5", 199.7, 206.5)]
    seq = _order_mixed(ln)
    check("the space goes where it is drawn",
          _reorder("".join(c for c, _, _ in ln), seq, len(ln)) if seq else None,
          "لخام عام 2015")
    check("letters drawn apart keep their space",
          _stray([("ا", 355, 358), ("م", 340, 345)], 0, 1, 206.6, 210.7), False)
    check("a space beside nothing is not moved",
          _stray([("ا", 355.2, 358.1), ("م", 350.5, 355.2)], 0, 1, 330, 334), False)
    from convert.order import _unjoined, JOINED
    check("the joining piece taken out of the text", _unjoined(f"الروس{JOINED}ي"), "الروسي")
    raw = {"blocks": [{"lines": [{"spans": [{"chars": [{"c": "س"}, {"c": JOINED}, {"c": "ي"}]}]}]}]}
    check("... and out of a detailed reading",
          [c["c"] for c in _unjoined(raw)["blocks"][0]["lines"][0]["spans"][0]["chars"]],
          ["س", "ي"])
    from convert import pdftext
    check("two-letter words reach the map", pdftext.MIN_FRAGMENT, 2)
    check("a two-letter word turned whole", apply_fixes("إن يب سي", {"يب": "بي"}), "إن بي سي")
    check("... never inside a longer word", apply_fixes("ترتيب", {"يب": "بي"}), "ترتيب")

    print(f"\n{len(FAILED)} failed" if FAILED else "\nall passed")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
